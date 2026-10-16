"""
Resume Parser — extracts text and sections from PDF, DOCX, or TXT uploads.
Uses PyMuPDF (layout-aware) for PDF parsing, python-docx for Word docs, and regex for section & bullet detection.
"""

import re
import io
import fitz  # PyMuPDF

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

from utils.constants import ACTION_VERBS  # Single source of truth


# Comprehensive section header patterns
SECTION_PATTERNS = {
    "summary": re.compile(
        r"^(professional\s+)?(summary|objective|profile|about(\s+me)?|overview|biography)$",
        re.IGNORECASE,
    ),
    "experience": re.compile(
        r"^(work\s+|professional\s+|relevant\s+|employment\s+)?(experience|history|background|employment|internships|work)$",
        re.IGNORECASE,
    ),
    "education": re.compile(
        r"^(education|academic|academics|qualifications?|degrees?|educational\s+background|academic\s+background)$",
        re.IGNORECASE,
    ),
    "skills": re.compile(
        r"^(technical\s+|core\s+|key\s+|professional\s+)?(skills|competencies|technologies|tools|stack|tech\s+stack)$",
        re.IGNORECASE,
    ),
    "projects": re.compile(
        r"^(key\s+|featured\s+|personal\s+|academic\s+|technical\s+)?(projects|portfolio|open\s+source)$",
        re.IGNORECASE,
    ),
    "certifications": re.compile(
        r"^(certifications?|certificates?|licenses?|awards?|achievements?|honors?)$",
        re.IGNORECASE,
    ),
    "leadership": re.compile(
        r"^(leadership|extracurricular(\s+activities)?|volunteer(\s+experience)?|positions\s+of\s+responsibility)$",
        re.IGNORECASE,
    ),
}


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extract raw text from PDF bytes using PyMuPDF.
    Sorts blocks topologically (top-to-bottom, left-to-right) for multi-column resumes.
    """
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    if len(doc) > 15:
        doc.close()
        raise ValueError("Resume exceeds maximum limit of 15 pages. Please upload a standard 1-3 page resume.")
    pages_text = []
    for page in doc:
        blocks = page.get_text("blocks", sort=True)
        page_lines = []
        for block in blocks:
            block_text = block[4].strip()
            if block_text:
                page_lines.append(block_text)
        pages_text.append("\n".join(page_lines))
    doc.close()
    return "\n\n".join(pages_text)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from Word .docx file bytes including tables."""
    if not DOCX_AVAILABLE:
        raise ImportError("python-docx is not installed. Please install it to parse .docx files.")

    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also extract any table content
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    return "\n".join(paragraphs)


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Decode plain text file bytes with encoding fallbacks."""
    for encoding in ("utf-8", "latin-1", "utf-16", "cp1252"):
        try:
            return file_bytes.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return file_bytes.decode("utf-8", errors="replace")


MAX_RESUME_SIZE = 5 * 1024 * 1024  # 5 MB limit to protect against OOM / DoS


def extract_text(uploaded_file) -> str:
    """
    Given a Streamlit UploadedFile, return extracted plain text.
    Supports .pdf, .docx, and .txt files. Enforces MAX_RESUME_SIZE guard.
    """
    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)
    file_bytes = uploaded_file.read(MAX_RESUME_SIZE + 1)
    if len(file_bytes) > MAX_RESUME_SIZE:
        raise ValueError(
            f"Uploaded file exceeds the maximum allowed size of {MAX_RESUME_SIZE // (1024 * 1024)} MB. "
            "Please upload a standard 1-3 page resume."
        )
    name = uploaded_file.name.lower()

    if name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif name.endswith(".txt"):
        return extract_text_from_txt(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {uploaded_file.name}. Please upload PDF, DOCX, or TXT.")


def detect_sections(text: str) -> dict[str, str]:
    """
    Split resume text into sections based on common header keywords.
    Uses strict header matching to prevent false positives within bullets.
    Returns a dict mapping section name -> section text.
    """
    lines = text.splitlines()
    sections: dict[str, list[str]] = {"other": []}
    current_section = "other"

    for line in lines:
        stripped = line.strip()
        if not stripped:
            # Preserve blank lines so section content retains paragraph spacing
            sections.setdefault(current_section, []).append("")
            continue

        # Clean punctuation from potential headers (e.g., "EXPERIENCE:" or "### Skills")
        clean_header = re.sub(r"^[\W_]+|[\W_]+$", "", stripped).strip()
        matched = False

        # Only evaluate lines that look like headers (short, < 50 chars, no terminal period)
        if 2 <= len(clean_header) <= 50 and not clean_header.endswith("."):
            for section_name, pattern in SECTION_PATTERNS.items():
                if pattern.match(clean_header):
                    current_section = section_name
                    if section_name not in sections:
                        sections[section_name] = []
                    matched = True
                    break

        if not matched:
            sections.setdefault(current_section, []).append(line)

    # Join each section's lines back into a string, strip leading/trailing whitespace
    return {k: "\n".join(v).strip() for k, v in sections.items() if any(l.strip() for l in v)}


def extract_bullet_points(text: str) -> list[str]:
    """
    Extract bullet point lines from resume text.
    Detects bullet symbols (•, -, *, ▪, ▫, –, —, >), numbered bullets (1., 2.),
    and action-oriented accomplishment lines.
    Filters out job titles, skill colon lists, education headers, and contact lines.
    """
    lines = text.splitlines()
    extracted_bullets = []

    # Symbol-prefixed bullets
    bullet_symbol_re = re.compile(
        r"^\s*[•\-\*\u2022\u25cf\u25e6\u25aa\u25ab\u25b8\u25ba\u2192\u2013\u2014>]+\s*(.+)$"
    )
    # Numbered bullets: 1. or 1)
    numbered_bullet_re = re.compile(r"^\s*\d+[\.\)]\s+(.+)$")
    # Date ranges common in job title headers (e.g. 2021 - 2024, Jan 2022 - Present)
    date_range_re = re.compile(
        r"(\b(19\d\d|20\d\d)\s*[-–—]\s*(19\d\d|20\d\d|present|current)\b|\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s+\d{4}\b)",
        re.IGNORECASE,
    )

    for line in lines:
        stripped = line.strip()
        if not stripped or len(stripped) < 12:  # lowered from 15 to catch short but valid bullets
            continue

        # Ignore lines that are clearly headings, contact info, or colon skill labels
        clean_no_punct = re.sub(r"^[\W_]+|[\W_]+$", "", stripped)
        if len(stripped) < 40 and any(p.match(clean_no_punct) for p in SECTION_PATTERNS.values()):
            continue
        if "@" in stripped or "linkedin.com" in stripped or "github.com" in stripped or "portfolio" in stripped.lower():
            continue
        # Ignore skill list headers like "Languages: Python, Go..." or "Tools & Frameworks: React, Docker"
        if re.match(r"^(languages|frameworks|tools|databases|libraries|technologies|stack|skills)\s*:", stripped, re.I):
            continue

        match_sym = bullet_symbol_re.match(stripped)
        if match_sym:
            b_text = match_sym.group(1).strip()
            if len(b_text) >= 12:
                extracted_bullets.append(b_text)
            continue

        match_num = numbered_bullet_re.match(stripped)
        if match_num:
            b_text = match_num.group(1).strip()
            if len(b_text) >= 12:
                extracted_bullets.append(b_text)
            continue

        # For un-prefixed lines: must NOT be a job title line with dates or company 'at'
        if date_range_re.search(stripped):
            continue
        if re.search(r"\b(at|@)\s+[A-Z][\w\s]{1,60}(?:\([0-9\s\-–—Present]+\))?$", stripped):
            continue

        # Check if the first word is a recognized action verb
        first_word = re.split(r"[\s\W]+", stripped)[0].lower()
        if first_word in ACTION_VERBS and len(stripped) >= 20:
            extracted_bullets.append(stripped)

    # Deduplicate while preserving order
    deduped = list(dict.fromkeys(extracted_bullets))
    return [b for b in deduped if len(b) >= 12]
